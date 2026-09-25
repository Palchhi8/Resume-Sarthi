"""Unit and integration test suite for AI Recruiter Agent pipeline."""

from datetime import date
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.parser import parse_resume, DocumentParsingError
from app.models import CandidateProfile, Experience, Education, HiringEvaluation
from app.workflow import (
    parse_date_string,
    calculate_employment_history_and_gaps,
    validate_candidate_profile,
    validate_hiring_evaluation,
)
from app.document_service import generate_evaluation_pdf
from app.email_service import dispatch_evaluation


@pytest.fixture
def client():
    """FastAPI TestClient fixture."""
    return TestClient(app)


# 1. Health Endpoint Test
def test_health_endpoint(client):
    """Verify health endpoint returns 200 and operational status."""
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert data["app"] == "AI Recruiter Agent"


# 2. TXT Parsing Test
def test_txt_parsing():
    """Verify standard text extraction works correctly."""
    sample_text = "John Doe\nSoftware Engineer\nPython, Docker\n"
    extracted = parse_resume(sample_text.encode("utf-8"), "resume.txt")
    assert "John Doe" in extracted
    assert "Python, Docker" in extracted


def test_pdf_parsing():
    """Verify PyMuPDF extracts text correctly from PDF bytes."""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Candidate Name: Morgan Stark\nTechnical Skills: Python, FastAPI, Docker")
    pdf_bytes = doc.tobytes()
    doc.close()

    extracted = parse_resume(pdf_bytes, "morgan.pdf")
    assert "Morgan Stark" in extracted
    assert "FastAPI" in extracted


def test_docx_parsing():
    """Verify python-docx extracts text correctly from DOCX bytes."""
    import docx
    import io
    doc = docx.Document()
    doc.add_paragraph("Candidate: Samantha Vance")
    doc.add_paragraph("Role: Cloud Architect")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extracted = parse_resume(docx_bytes, "samantha.docx")
    assert "Samantha Vance" in extracted
    assert "Cloud Architect" in extracted



# 3. Unsupported File Handling Test
def test_unsupported_file_handling():
    """Verify unsupported file extensions raise DocumentParsingError."""
    with pytest.raises(DocumentParsingError) as exc_info:
        parse_resume(b"Fake audio content", "candidate_audio.mp3")
    assert "Unsupported file format" in str(exc_info.value)


# 4. Empty Resume Handling Test
def test_empty_resume_handling():
    """Verify empty byte files or blank documents raise DocumentParsingError."""
    with pytest.raises(DocumentParsingError) as exc_info:
        parse_resume(b"", "empty_resume.txt")
    assert "Unable to extract readable text" in str(exc_info.value)

    with pytest.raises(DocumentParsingError) as exc_info:
        parse_resume(b"   \n\t  ", "blank_resume.txt")
    assert "Unable to extract readable text" in str(exc_info.value)


# 5. Candidate Schema Validation Test
def test_candidate_schema_validation():
    """Verify CandidateProfile and nested models validate and serialize properly."""
    exp = Experience(
        company="Tech Corp",
        role="Backend Dev",
        start_date="June 2020",
        end_date="December 2021",
        technologies=["Python", "FastAPI"],
    )
    edu = Education(
        degree="B.S. Computer Science",
        institution="State University",
        year="2020",
    )
    candidate = CandidateProfile(
        name="Alex Smith",
        email="alex.smith@example.com",
        skills=["Python", "AWS", "Docker"],
        experience=[exp],
        education=[edu],
    )
    assert candidate.name == "Alex Smith"
    assert len(candidate.experience) == 1
    assert candidate.experience[0].technologies == ["Python", "FastAPI"]
    assert candidate.years_of_experience is None


# 6. Employment Gap Calculation Test (Deterministic)
def test_employment_gap_calculation_detected():
    """
    Verify deterministic gap detection:
    June 2020 - December 2021 (Acme)
    April 2022 - Present (Beta)
    Should detect ~4-month gap between Dec 2021 and Apr 2022.
    """
    candidate = CandidateProfile(
        name="Taylor Reed",
        experience=[
            Experience(
                company="Beta Corp",
                role="Senior Engineer",
                start_date="April 2022",
                end_date="Present",
                is_current=True,
            ),
            Experience(
                company="Acme Inc",
                role="Software Engineer",
                start_date="June 2020",
                end_date="December 2021",
                is_current=False,
            ),
        ],
    )
    analyzed = calculate_employment_history_and_gaps(candidate)
    assert len(analyzed.employment_gaps) == 1
    gap_msg = analyzed.employment_gaps[0]
    assert "December 2021" in gap_msg
    assert "April 2022" in gap_msg
    assert "gap" in gap_msg.lower()
    assert analyzed.years_of_experience is not None
    assert analyzed.years_of_experience > 2.0


def test_employment_gap_overlapping_roles():
    """Verify overlapping roles do not distort true gap between employment end and next start."""
    candidate = CandidateProfile(
        name="Jordan Tech",
        experience=[
            Experience(
                company="Primary Corp",
                role="Full-time Engineer",
                start_date="January 2020",
                end_date="December 2021",
            ),
            Experience(
                company="Consulting Gig",
                role="Advisor",
                start_date="March 2021",
                end_date="August 2021",
            ),
            Experience(
                company="Next Corp",
                role="Lead Dev",
                start_date="April 2022",
                end_date="Present",
                is_current=True,
            ),
        ],
    )
    analyzed = calculate_employment_history_and_gaps(candidate)
    assert len(analyzed.employment_gaps) == 1
    gap_msg = analyzed.employment_gaps[0]
    assert "December 2021" in gap_msg
    assert "April 2022" in gap_msg
    assert "~4-month" in gap_msg



def test_employment_gap_continuous_employment():
    """Verify no gaps are flagged when candidate transitions directly with <= 1 month."""
    candidate = CandidateProfile(
        name="Morgan Lee",
        experience=[
            Experience(
                company="First Corp",
                role="Junior Dev",
                start_date="January 2020",
                end_date="December 2020",
            ),
            Experience(
                company="Second Corp",
                role="Mid Dev",
                start_date="January 2021",
                end_date="December 2021",
            ),
        ],
    )
    analyzed = calculate_employment_history_and_gaps(candidate)
    assert len(analyzed.employment_gaps) == 0


def test_employment_gap_ambiguous_dates():
    """Verify ambiguous or unparseable dates result in explicit non-fabrication ambiguity note."""
    candidate = CandidateProfile(
        name="Casey Brown",
        experience=[
            Experience(
                company="Unknown Entity",
                role="Consultant",
                start_date="A few years ago",
                end_date="Recently",
            ),
        ],
    )
    analyzed = calculate_employment_history_and_gaps(candidate)
    assert any("Employment gap cannot be reliably determined" in amb for amb in analyzed.ambiguities)


# 7. Candidate Profile Validation Test
def test_candidate_profile_validation():
    """Verify missing phone, email, and ambiguous titles are surfaced."""
    candidate = CandidateProfile(
        name="Sam Miller",
        email=None,  # Missing email
        phone=None,  # Missing phone
        experience=[
            Experience(
                company="Self",
                role="Freelance Consultant / Specialist",
                start_date="2021",
                end_date="2022",
            )
        ],
    )
    validated = validate_candidate_profile(candidate)
    assert any("Missing candidate contact email" in amb for amb in validated.ambiguities)
    assert any("Missing direct phone number" in amb for amb in validated.ambiguities)
    assert any("ambiguous" in amb.lower() for amb in validated.ambiguities)


# 8. Evaluation Schema Validation Test
def test_evaluation_schema_validation():
    """Verify HiringEvaluation model and validation function."""
    eval_model = HiringEvaluation(
        candidate_name="Dana Scully",
        email="dana@fbi.mock",
        primary_skillset=["Forensics", "Python", "Data Analysis"],
        years_of_experience=5.0,
        education=["M.D. Forensic Pathology"],
        relevant_experience=["Senior Field Investigator"],
        employment_gaps=[],
        ambiguities_missing_info=[],
        recommended_role="Senior Forensic Engineer",
        evaluation_summary="Candidate demonstrates rigorous analytical capabilities and strong track record.",
    )
    validated = validate_hiring_evaluation(eval_model)
    assert validated.candidate_name == "Dana Scully"
    assert validated.recommended_role == "Senior Forensic Engineer"
    assert validated.decision_recommendation == "Proceed to Technical Interview"


# 9. Document PDF Generation Test
def test_pdf_generation(tmp_path):
    """Verify ReportLab produces valid PDF file without crashing."""
    eval_model = HiringEvaluation(
        candidate_name="Test Candidate",
        email="test@company.mock",
        primary_skillset=["FastAPI", "Docker", "AWS"],
        years_of_experience=3.5,
        education=["B.S. CS - University (2020)"],
        relevant_experience=["Built high performance backend APIs in Python"],
        employment_gaps=["Detected ~4-month gap between Dec 2021 and Apr 2022"],
        ambiguities_missing_info=["Missing candidate phone number"],
        recommended_role="Backend Software Engineer",
        evaluation_summary="Thorough candidate with demonstrated skills in Python backend engineering.",
        strengths=["API development", "AWS containerization"],
        areas_for_investigation=["Clarify 4-month gap in 2022"],
    )
    pdf_path, filename = generate_evaluation_pdf(eval_model)
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 1000
    assert filename.endswith(".pdf")


# 10. Mock Email Dispatch Test
def test_mock_email_dispatch():
    """Verify mock email dispatch returns structured simulation response."""
    # Create sample evaluation to dispatch
    eval_model = HiringEvaluation(
        candidate_name="Dispatch Test Candidate",
        recommended_role="Staff Engineer",
        evaluation_summary="Standard evaluation for dispatch test.",
    )
    pdf_path, filename = generate_evaluation_pdf(eval_model)

    response = dispatch_evaluation(file_path=str(pdf_path), recipient="hr-admissions@company.mock")
    assert response.status == "simulated"
    assert response.recipient == "hr-admissions@company.mock"
    assert response.attachment == filename
    assert "succeeded" in response.message.lower()
