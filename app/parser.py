"""Document parsing module for extracting text from PDF, DOCX, and TXT files."""

import io
from pathlib import Path
import docx
import pymupdf


class DocumentParsingError(Exception):
    """Custom exception raised when document parsing fails."""
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes using PyMuPDF."""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        pages_text = []
        for page_index in range(len(doc)):
            page = doc[page_index]
            text = page.get_text("text")
            if text:
                pages_text.append(text.strip())
        doc.close()
        return "\n\n".join(pages_text).strip()
    except Exception as exc:
        raise DocumentParsingError(f"Failed to read PDF document: {exc}") from exc


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_texts.append(row_text)
                    
        combined = paragraphs + table_texts
        return "\n".join(combined).strip()
    except Exception as exc:
        raise DocumentParsingError(f"Failed to read DOCX document: {exc}") from exc


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from plain text bytes."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise DocumentParsingError("Unable to decode text file with standard encodings.")


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """
    Parse uploaded resume file bytes and return clean readable text.
    
    Supported extensions: .pdf, .docx, .txt
    Raises DocumentParsingError with user-friendly messages for unsupported
    formats or unreadable/empty documents.
    """
    if not file_bytes:
        raise DocumentParsingError("Unable to extract readable text from this document. File is empty.")

    suffix = Path(filename).suffix.lower()
    
    if suffix == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    elif suffix == ".docx":
        text = extract_text_from_docx(file_bytes)
    elif suffix == ".txt":
        text = extract_text_from_txt(file_bytes)
    else:
        raise DocumentParsingError(
            f"Unsupported file format '{suffix}'. Please upload a PDF (.pdf), Word document (.docx), or plain text (.txt) file."
        )

    clean_text = text.strip()
    if not clean_text:
        raise DocumentParsingError("Unable to extract readable text from this document.")

    return clean_text
